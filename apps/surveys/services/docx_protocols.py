from io import BytesIO

from django.utils.timezone import localtime
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from docx.shared import Pt

from apps.surveys.services.results import get_survey_results

from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Max

from apps.surveys.models import GeneratedProtocol


MONTHS_GENITIVE = {
    1: "января",
    2: "февраля",
    3: "марта",
    4: "апреля",
    5: "мая",
    6: "июня",
    7: "июля",
    8: "августа",
    9: "сентября",
    10: "октября",
    11: "ноября",
    12: "декабря",
}


def format_protocol_date(value):
    """
    Формат:
    «21» мая 2026 г.
    """

    if value is None:
        return ""

    month_name = MONTHS_GENITIVE[value.month]

    return (
        f"«{value.day:02d}» "
        f"{month_name} "
        f"{value.year} г."
    )


def set_cell_text(
    cell,
    text,
    *,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    font_size=8,
):
    """
    Заполняет ячейку таблицы единообразным текстом.
    """

    cell.text = ""

    paragraph = cell.paragraphs[0]
    paragraph.alignment = align

    run = paragraph.add_run(
        "" if text is None else str(text)
    )

    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(font_size)

    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )


def set_document_default_font(document):
    """
    Устанавливает Arial 9 pt для базового стиля.
    """

    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(9)


def add_text_paragraph(
    document,
    label,
    value,
):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)

    label_run = paragraph.add_run(label)
    label_run.bold = True
    label_run.font.name = "Arial"
    label_run.font.size = Pt(9)

    value_run = paragraph.add_run(
        "" if value is None else str(value)
    )

    value_run.font.name = "Arial"
    value_run.font.size = Pt(9)

    return paragraph


def get_protocol_participants(survey):
    """
    Возвращает участников, учитываемых в расчётах.

    Формат:
    Иванов И.И. – инженер-технолог
    """

    submissions = (
        survey.submissions
        .filter(
            is_excluded=False,
        )
        .order_by(
            "submitted_at",
            "id",
        )
    )

    participants = []
    seen = set()

    for submission in submissions:
        key = (
            submission.full_name.strip().casefold(),
            submission.position.strip().casefold(),
        )

        if key in seen:
            continue

        seen.add(key)

        if submission.position:
            participants.append(
                f"{submission.full_name} – "
                f"{submission.position}"
            )
        else:
            participants.append(
                submission.full_name
            )

    return participants


def build_protocol_docx(
    *,
    survey,
    protocol,
):
    """
    Формирует DOCX-протокол и возвращает его содержимое bytes.
    """

    results = get_survey_results(survey)
    participants = get_protocol_participants(survey)

    document = Document()
    set_document_default_font(document)

    section = document.sections[0]

    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)

    section.top_margin = Cm(1)
    section.bottom_margin = Cm(1)
    section.left_margin = Cm(1)
    section.right_margin = Cm(1)

    if protocol.document_code:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_after = Pt(2)

        run = paragraph.add_run(protocol.document_code)

        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)

    title_run = title.add_run(
        f"ПРОТОКОЛ дегустации № "
        f"{protocol.protocol_number} от "
        f"{format_protocol_date(protocol.protocol_date)}"
    )

    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(12)

    add_text_paragraph(
        document,
        "Сотрудник отдела разработки продуктов: ",
        protocol.responsible_employee,
    )

    add_text_paragraph(
        document,
        "Присутствовали: ",
        "; ".join(participants),
    )

    add_text_paragraph(
        document,
        "Количество участников: ",
        (
            f"{results['submissions_included']} "
            f"участников."
        ),
    )

    add_text_paragraph(
        document,
        "Цель дегустации: ",
        protocol.tasting_goal,
    )

    if protocol.room_conditions:
        paragraph = document.add_paragraph(
            protocol.room_conditions
        )
        paragraph.paragraph_format.space_after = Pt(2)

    if protocol.product_conditions:
        paragraph = document.add_paragraph(
            protocol.product_conditions
        )
        paragraph.paragraph_format.space_after = Pt(4)

    rating_questions = results[
        "rating_questions"
    ]

    column_count = (
        3
        + len(rating_questions)
        + 2
    )

    table = document.add_table(
        rows=2,
        cols=column_count,
    )

    table.style = "Table Grid"
    table.autofit = False

    first_header_row = table.rows[0].cells
    second_header_row = table.rows[1].cells

    set_cell_text(
        first_header_row[0],
        "№\nп/п",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    set_cell_text(
        first_header_row[1],
        "Название продукции",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    set_cell_text(
        first_header_row[2],
        "Характеристика продукта",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    first_header_row[0].merge(
        second_header_row[0]
    )

    first_header_row[1].merge(
        second_header_row[1]
    )

    first_header_row[2].merge(
        second_header_row[2]
    )

    rating_start = 3
    rating_end = (
        rating_start
        + len(rating_questions)
        - 1
    )

    average_column = rating_end + 1
    comments_column = average_column + 1

    if rating_questions:
        rating_header = first_header_row[
            rating_start
        ]

        rating_header.merge(
            first_header_row[rating_end]
        )

        set_cell_text(
            rating_header,
            "Балловая оценка "
            "(по 5-ти балльной шкале)",
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for index, question in enumerate(
        rating_questions,
        start=rating_start,
    ):
        set_cell_text(
            second_header_row[index],
            question.title,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    set_cell_text(
        first_header_row[average_column],
        "Средний балл",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    first_header_row[average_column].merge(
        second_header_row[average_column]
    )

    set_cell_text(
        first_header_row[comments_column],
        "Комментарии к балловым оценкам",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    first_header_row[comments_column].merge(
        second_header_row[comments_column]
    )

    for row_number, row_data in enumerate(
        results["rows"],
        start=1,
    ):
        row_cells = table.add_row().cells

        set_cell_text(
            row_cells[0],
            row_number,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

        set_cell_text(
            row_cells[1],
            row_data["sample"].name,
        )

        set_cell_text(
            row_cells[2],
            row_data["sample"].description,
        )

        for index, score_data in enumerate(
            row_data["scores"],
            start=rating_start,
        ):
            score = score_data["average"]

            set_cell_text(
                row_cells[index],
                "—" if score is None else score,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )

        set_cell_text(
            row_cells[average_column],
            (
                "—"
                if row_data["overall_average"] is None
                else row_data["overall_average"]
            ),
            align=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
        )

        set_cell_text(
            row_cells[comments_column],
            "; ".join(row_data["comments"]),
        )

    widths = {
        0: Cm(1),
        1: Cm(3),
        2: Cm(6.5),
        comments_column: Cm(5.5),
    }

    for column_index in range(
        rating_start,
        average_column + 1,
    ):
        widths[column_index] = Cm(1.5)

    for row in table.rows:
        for column_index, width in widths.items():
            row.cells[column_index].width = width

    document.add_paragraph()

    add_text_paragraph(
        document,
        "Заключение: ",
        protocol.conclusion,
    )

    document.add_paragraph()

    signature = document.add_paragraph()
    signature.paragraph_format.space_before = Pt(8)

    signature_run = signature.add_run(
        f"{protocol.signer_position}"
        f"______________/ "
        f"{protocol.signer_name}"
    )

    signature_run.font.name = "Arial"
    signature_run.font.size = Pt(9)

    footer = document.sections[0].footer

    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    footer_run = footer_paragraph.add_run(
        "Сформировано системой «Леовит формы»"
    )

    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(7)

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def sanitize_filename_part(value):
    """
    Удаляет символы, нежелательные в имени файла.
    """

    safe_value = "".join(
        character
        for character in str(value)
        if (
            character.isalnum()
            or character in {"-", "_"}
        )
    )

    return safe_value or "protocol"


@transaction.atomic
def generate_and_save_protocol(
    *,
    survey,
    protocol,
    user,
):
    """
    Создаёт новую версию протокола и сохраняет DOCX.
    """

    maximum_version = (
        survey.generated_protocols.aggregate(
            maximum=Max("version"),
        )["maximum"]
        or 0
    )

    version = maximum_version + 1

    generated_protocol = GeneratedProtocol(
        survey=survey,
        protocol_number=protocol.protocol_number,
        version=version,
        generated_by=user,
    )

    document_data = build_protocol_docx(
        survey=survey,
        protocol=protocol,
    )

    safe_number = sanitize_filename_part(
        protocol.protocol_number
    )

    filename = (
        f"protocol_{safe_number}_"
        f"v{version}.docx"
    )

    generated_protocol.file.save(
        filename,
        ContentFile(document_data),
        save=False,
    )

    generated_protocol.save()

    return generated_protocol

