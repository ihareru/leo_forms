import tempfile
from datetime import date
from decimal import Decimal
from io import BytesIO
from docx import Document
from django.urls import reverse
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.storage import default_storage
from django.test import TestCase, override_settings

from .forms import RatingDecimalField, PublicSurveyForm, SurveyProtocolForm
from .models import (
    Survey,
    SurveyQuestion,
    SurveySample,
    Answer,
    Submission,
    GeneratedProtocol,
    SurveyProtocol,
)
from .services.results import get_survey_results
from .services.docx_protocols import build_protocol_docx, generate_and_save_protocol
from .services.protocol_defaults import build_protocol_initial_data, get_next_protocol_number
from .services.copying import copy_survey


User = get_user_model()


class SurveyModelTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="creator",
            password="StrongPassword123!",
        )

    def test_new_survey_is_draft(self):
        survey = Survey.objects.create(
            owner=self.user,
            title="Тестовая форма",
        )

        self.assertEqual(
            survey.status,
            Survey.Status.DRAFT,
        )

        self.assertFalse(
            survey.accepts_responses,
        )

    def test_default_questions_created(self):
        survey = Survey.objects.create(
            owner=self.user,
            title="Тестовая форма",
        )

        self.assertEqual(
            survey.questions.count(),
            6,
        )

        question_codes = set(
            survey.questions.values_list(
                "code",
                flat=True,
            )
        )

        expected_codes = {
            SurveyQuestion.StandardCode.APPEARANCE,
            SurveyQuestion.StandardCode.COLOR,
            SurveyQuestion.StandardCode.CONSISTENCY,
            SurveyQuestion.StandardCode.SMELL,
            SurveyQuestion.StandardCode.TASTE,
            SurveyQuestion.StandardCode.COMMENT,
        }

        self.assertEqual(
            question_codes,
            expected_codes,
        )

    def test_public_id_is_unique(self):
        survey_1 = Survey.objects.create(
            owner=self.user,
            title="Форма №1",
        )

        survey_2 = Survey.objects.create(
            owner=self.user,
            title="Форма №2",
        )

        self.assertNotEqual(
            survey_1.public_id,
            survey_2.public_id,
        )

    def test_cannot_publish_survey_without_samples(self):
        survey = Survey.objects.create(
            owner=self.user,
            title="Форма без образцов",
        )

        with self.assertRaises(ValidationError):
            survey.publish()

    def test_survey_can_be_published_with_sample(self):
        survey = Survey.objects.create(
            owner=self.user,
            title="Форма с образцом",
        )

        SurveySample.objects.create(
            survey=survey,
            name="Образец №1",
            order=10,
        )

        survey.publish()
        survey.refresh_from_db()

        self.assertEqual(
            survey.status,
            Survey.Status.PUBLISHED,
        )

        self.assertTrue(
            survey.accepts_responses,
        )

        self.assertIsNotNone(
            survey.published_at,
        )

    def test_published_survey_can_be_closed(self):
        survey = Survey.objects.create(
            owner=self.user,
            title="Форма с образцом",
        )

        SurveySample.objects.create(
            survey=survey,
            name="Образец №1",
            order=10,
        )

        survey.publish()
        survey.close()
        survey.refresh_from_db()

        self.assertEqual(
            survey.status,
            Survey.Status.CLOSED,
        )

        self.assertFalse(
            survey.accepts_responses,
        )

        self.assertIsNotNone(
            survey.closed_at,
        )

    def test_draft_survey_cannot_be_closed(self):
        survey = Survey.objects.create(
            owner=self.user,
            title="Черновик",
        )

        with self.assertRaises(ValidationError):
            survey.close()


class RatingDecimalFieldTests(TestCase):
    def setUp(self):
        self.field = RatingDecimalField()

    def test_accepts_decimal_with_dot(self):
        value = self.field.clean("3.5")

        self.assertEqual(
            value,
            Decimal("3.5"),
        )

    def test_accepts_decimal_with_comma(self):
        value = self.field.clean("3,5")

        self.assertEqual(
            value,
            Decimal("3.5"),
        )

    def test_accepts_integer(self):
        value = self.field.clean("5")

        self.assertEqual(
            value,
            Decimal("5"),
        )

    def test_rejects_value_below_minimum(self):
        with self.assertRaises(ValidationError):
            self.field.clean("0,9")

    def test_rejects_value_above_maximum(self):
        with self.assertRaises(ValidationError):
            self.field.clean("5,1")

    def test_rejects_non_numeric_value(self):
        with self.assertRaises(ValidationError):
            self.field.clean("пять")


class PublicSurveyFormTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="creator_public",
            password="StrongPassword123!",
        )

        self.survey = Survey.objects.create(
            owner=self.user,
            title="Публичная форма",
        )

        self.sample = SurveySample.objects.create(
            survey=self.survey,
            name="Образец №1",
            order=10,
        )

    def get_valid_form_data(self):
        data = {
            "full_name": "Иванов Иван Иванович",
            "position": "Инженер-технолог",
        }

        for question in self.survey.questions.all():
            field_name = (
                PublicSurveyForm.get_answer_field_name(
                    sample_id=self.sample.pk,
                    question_id=question.pk,
                )
            )

            if (
                question.question_type
                == SurveyQuestion.QuestionType.RATING
            ):
                data[field_name] = "4,5"
            else:
                data[field_name] = "Хороший образец"

        return data

    def test_public_form_accepts_comma_scores(self):
        form = PublicSurveyForm(
            data=self.get_valid_form_data(),
            survey=self.survey,
        )

        self.assertTrue(
            form.is_valid(),
            form.errors,
        )

        rating_question = self.survey.questions.filter(
            question_type=(
                SurveyQuestion.QuestionType.RATING
            ),
        ).first()

        field_name = form.get_answer_field_name(
            sample_id=self.sample.pk,
            question_id=rating_question.pk,
        )

        self.assertEqual(
            form.cleaned_data[field_name],
            Decimal("4.5"),
        )

    def test_public_form_rejects_score_above_five(self):
        data = self.get_valid_form_data()

        rating_question = self.survey.questions.filter(
            question_type=(
                SurveyQuestion.QuestionType.RATING
            ),
        ).first()

        field_name = (
            PublicSurveyForm.get_answer_field_name(
                sample_id=self.sample.pk,
                question_id=rating_question.pk,
            )
        )

        data[field_name] = "5,1"

        form = PublicSurveyForm(
            data=data,
            survey=self.survey,
        )

        self.assertFalse(form.is_valid())

        self.assertIn(
            field_name,
            form.errors,
        )

    def get_comment_question(self):
        return self.survey.questions.get(
            question_type=(SurveyQuestion.QuestionType.TEXT),
        )

    def set_all_rating_values(
        self,
        data,
        value,
    ):
        for question in self.survey.questions.filter(
            question_type=(SurveyQuestion.QuestionType.RATING),
        ):
            field_name = PublicSurveyForm.get_answer_field_name(
                sample_id=self.sample.pk,
                question_id=question.pk,
            )

            data[field_name] = value

    def test_comment_required_when_score_below_five(self):
        data = self.get_valid_form_data()

        comment_question = self.get_comment_question()

        comment_field_name = PublicSurveyForm.get_answer_field_name(
            sample_id=self.sample.pk,
            question_id=comment_question.pk,
        )

        data[comment_field_name] = ""

        form = PublicSurveyForm(
            data=data,
            survey=self.survey,
        )

        self.assertFalse(
            form.is_valid(),
        )

        self.assertIn(
            comment_field_name,
            form.errors,
        )

        self.assertIn(
            "При оценке ниже 5,0 необходимо указать комментарий.",
            form.errors[comment_field_name],
        )

    def test_comment_not_required_when_all_scores_are_five(self):
        data = self.get_valid_form_data()

        self.set_all_rating_values(
            data,
            "5,0",
        )

        comment_question = self.get_comment_question()

        comment_field_name = PublicSurveyForm.get_answer_field_name(
            sample_id=self.sample.pk,
            question_id=comment_question.pk,
        )

        data[comment_field_name] = ""

        form = PublicSurveyForm(
            data=data,
            survey=self.survey,
        )

        self.assertTrue(
            form.is_valid(),
            form.errors,
        )

    def test_comment_accepts_score_below_five_when_filled(self):
        data = self.get_valid_form_data()

        comment_question = self.get_comment_question()

        comment_field_name = PublicSurveyForm.get_answer_field_name(
            sample_id=self.sample.pk,
            question_id=comment_question.pk,
        )

        data[comment_field_name] = "Недостаточно выраженный вкус."

        form = PublicSurveyForm(
            data=data,
            survey=self.survey,
        )

        self.assertTrue(
            form.is_valid(),
            form.errors,
        )


class PublicSurveyViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="public_owner",
            password="StrongPassword123!",
        )

        self.survey = Survey.objects.create(
            owner=self.user,
            title="Дегустация",
        )

        self.sample = SurveySample.objects.create(
            survey=self.survey,
            name="Образец №1",
            order=10,
        )

    def get_public_url(self):
        return reverse(
            "public_survey",
            args=[
                self.survey.public_id,
            ],
        )

    def get_valid_submission_data(self):
        data = {
            "full_name": "Петров Пётр Петрович",
            "position": "Технолог",
        }

        for question in self.survey.questions.all():
            field_name = (
                PublicSurveyForm.get_answer_field_name(
                    sample_id=self.sample.pk,
                    question_id=question.pk,
                )
            )

            if (
                question.question_type
                == SurveyQuestion.QuestionType.RATING
            ):
                data[field_name] = "3,5"
            else:
                data[field_name] = (
                    "Комментарий участника"
                )

        return data

    def test_draft_survey_is_not_public(self):
        response = self.client.get(
            self.get_public_url(),
        )

        self.assertEqual(
            response.status_code,
            404,
        )

    def test_published_survey_is_public(self):
        self.survey.publish()

        response = self.client.get(
            self.get_public_url(),
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertContains(
            response,
            self.survey.title,
        )

        self.assertContains(
            response,
            self.sample.name,
        )

    def test_closed_survey_does_not_accept_answers(self):
        self.survey.publish()
        self.survey.close()

        response = self.client.get(
            self.get_public_url(),
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertContains(
            response,
            "Сбор ответов завершён",
        )

    def test_submission_is_saved(self):
        self.survey.publish()

        response = self.client.post(
            self.get_public_url(),
            self.get_valid_submission_data(),
        )

        self.assertEqual(
            Submission.objects.count(),
            1,
        )

        submission = Submission.objects.get()

        self.assertEqual(
            submission.full_name,
            "Петров Пётр Петрович",
        )

        self.assertEqual(
            submission.position,
            "Технолог",
        )

        expected_answer_count = (
            self.survey.questions.count()
            * self.survey.samples.filter(
                is_active=True,
            ).count()
        )

        self.assertEqual(
            submission.answers.count(),
            expected_answer_count,
        )

        self.assertEqual(
            response.status_code,
            302,
        )

    def test_comma_score_saved_as_decimal(self):
        self.survey.publish()

        self.client.post(
            self.get_public_url(),
            self.get_valid_submission_data(),
        )

        answer = Answer.objects.filter(
            question__question_type=(
                SurveyQuestion.QuestionType.RATING
            ),
        ).first()

        self.assertEqual(
            answer.numeric_value,
            Decimal("3.5"),
        )

    def test_comment_is_linked_to_participant(self):
        self.survey.publish()

        self.client.post(
            self.get_public_url(),
            self.get_valid_submission_data(),
        )

        comment_answer = Answer.objects.get(
            question__question_type=(
                SurveyQuestion.QuestionType.TEXT
            ),
        )

        self.assertEqual(
            comment_answer.text_value,
            "Комментарий участника",
        )

        self.assertEqual(
            comment_answer.submission.full_name,
            "Петров Пётр Петрович",
        )

    def test_second_submission_blocked_in_same_session(self):
        self.survey.publish()

        data = self.get_valid_submission_data()

        self.client.post(
            self.get_public_url(),
            data,
        )

        response = self.client.post(
            self.get_public_url(),
            data,
        )

        self.assertEqual(
            Submission.objects.count(),
            1,
        )

        self.assertContains(
            response,
            "Ответ уже отправлен",
        )

    def test_multiple_submissions_can_be_allowed(self):
        self.survey.allow_multiple_submissions = True
        self.survey.save(
            update_fields=[
                "allow_multiple_submissions",
                "updated_at",
            ]
        )

        self.survey.publish()

        data = self.get_valid_submission_data()

        self.client.post(
            self.get_public_url(),
            data,
        )

        self.client.post(
            self.get_public_url(),
            data,
        )

        self.assertEqual(
            Submission.objects.count(),
            2,
        )

    def test_inactive_sample_is_not_shown(self):
        self.sample.is_active = False
        self.sample.save(
            update_fields=[
                "is_active",
                "updated_at",
            ]
        )

        active_sample = SurveySample.objects.create(
            survey=self.survey,
            name="Активный образец",
            order=20,
        )

        self.survey.publish()

        response = self.client.get(
            self.get_public_url(),
        )

        self.assertNotContains(
            response,
            self.sample.name,
        )

        self.assertContains(
            response,
            active_sample.name,
        )


class SurveyResultsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="results_owner",
            password="StrongPassword123!",
        )

        self.survey = Survey.objects.create(
            owner=self.user,
            title="Результаты дегустации",
        )

        self.sample = SurveySample.objects.create(
            survey=self.survey,
            name="Образец №1",
            order=10,
        )

        self.rating_question = (
            self.survey.questions.filter(
                question_type=(
                    SurveyQuestion.QuestionType.RATING
                ),
            )
            .order_by(
                "order",
            )
            .first()
        )

        self.comment_question = (
            self.survey.questions.filter(
                question_type=(
                    SurveyQuestion.QuestionType.TEXT
                ),
            )
            .first()
        )

    def create_submission(
        self,
        *,
        full_name,
        score,
        comment="",
        is_excluded=False,
    ):
        submission = Submission.objects.create(
            survey=self.survey,
            full_name=full_name,
            position="Технолог",
            is_excluded=is_excluded,
        )

        Answer.objects.create(
            submission=submission,
            sample=self.sample,
            question=self.rating_question,
            numeric_value=score,
        )

        Answer.objects.create(
            submission=submission,
            sample=self.sample,
            question=self.comment_question,
            text_value=comment,
        )

        return submission

    def test_average_score_is_calculated(self):
        self.create_submission(
            full_name="Участник №1",
            score=Decimal("4.0"),
        )

        self.create_submission(
            full_name="Участник №2",
            score=Decimal("5.0"),
        )

        results = get_survey_results(
            self.survey,
        )

        rating_score = results["rows"][0]["scores"][0]

        self.assertEqual(
            rating_score["average"],
            Decimal("4.5"),
        )

    def test_excluded_submission_is_not_in_average(self):
        self.create_submission(
            full_name="Участник №1",
            score=Decimal("5.0"),
        )

        self.create_submission(
            full_name="Ошибочный участник",
            score=Decimal("1.0"),
            is_excluded=True,
        )

        results = get_survey_results(
            self.survey,
        )

        rating_score = results["rows"][0]["scores"][0]

        self.assertEqual(
            rating_score["average"],
            Decimal("5.0"),
        )

        self.assertEqual(
            results["submissions_total"],
            2,
        )

        self.assertEqual(
            results["submissions_included"],
            1,
        )

        self.assertEqual(
            results["submissions_excluded"],
            1,
        )

    def test_comments_are_anonymous_in_results(self):
        self.create_submission(
            full_name="Иванов Иван Иванович",
            score=Decimal("4.0"),
            comment="Недостаточно выраженный вкус",
        )

        results = get_survey_results(
            self.survey,
        )

        comments = results["rows"][0]["comments"]

        self.assertEqual(
            comments,
            [
                "Недостаточно выраженный вкус",
            ],
        )

        self.assertNotIn(
            "Иванов Иван Иванович",
            comments,
        )

    def test_overall_average_uses_two_decimal_places(self):
        submission = Submission.objects.create(
            survey=self.survey,
            full_name="Участник",
            position="Технолог",
        )

        rating_questions = list(
            self.survey.questions.filter(
                question_type=(SurveyQuestion.QuestionType.RATING),
            ).order_by(
                "order",
                "id",
            )
        )

        scores = [
            Decimal("5.0"),
            Decimal("5.0"),
            Decimal("5.0"),
            Decimal("5.0"),
            Decimal("4.9"),
        ]

        for question, score in zip(
            rating_questions,
            scores,
        ):
            Answer.objects.create(
                submission=submission,
                sample=self.sample,
                question=question,
                numeric_value=score,
            )

        results = get_survey_results(
            self.survey,
        )

        self.assertEqual(
            results["rows"][0]["overall_average"],
            Decimal("4.98"),
        )

    def test_question_average_uses_two_decimal_places(self):
        self.create_submission(
            full_name="Участник №1",
            score=Decimal("4.9"),
        )

        self.create_submission(
            full_name="Участник №2",
            score=Decimal("5.0"),
        )

        results = get_survey_results(
            self.survey,
        )

        rating_score = results["rows"][0]["scores"][0]

        self.assertEqual(
            rating_score["average"],
            Decimal("4.95"),
        )

class SurveyReportViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="report_owner",
            password="StrongPassword123!",
        )

        self.other_user = User.objects.create_user(
            username="report_other",
            password="StrongPassword123!",
        )

        self.survey = Survey.objects.create(
            owner=self.user,
            title="Форма для отчётов",
        )

        self.sample = SurveySample.objects.create(
            survey=self.survey,
            name="Образец №1",
            order=10,
        )

    def test_owner_can_open_results(self):
        self.client.force_login(self.user)

        response = self.client.get(
            reverse(
                "surveys:survey_results",
                args=[
                    self.survey.pk,
                ],
            ),
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertContains(
            response,
            "Результаты",
        )

    def test_other_user_cannot_open_results(self):
        self.client.force_login(self.other_user)

        response = self.client.get(
            reverse(
                "surveys:survey_results",
                args=[
                    self.survey.pk,
                ],
            ),
        )

        self.assertEqual(
            response.status_code,
            404,
        )

    def test_qr_code_is_png(self):
        self.client.force_login(self.user)

        response = self.client.get(
            reverse(
                "surveys:survey_qr_code",
                args=[
                    self.survey.pk,
                ],
            ),
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertEqual(
            response["Content-Type"],
            "image/png",
        )

        self.assertTrue(
            response.content.startswith(
                b"\x89PNG",
            )
        )

    def test_excel_export_returns_xlsx(self):
        self.client.force_login(self.user)

        response = self.client.get(
            reverse(
                "surveys:survey_excel_export",
                args=[
                    self.survey.pk,
                ],
            ),
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertEqual(
            response["Content-Type"],
            (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

        self.assertTrue(
            response.content.startswith(
                b"PK",
            )
        )

class SubmissionManagementTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="submission_owner",
            password="StrongPassword123!",
        )

        self.survey = Survey.objects.create(
            owner=self.user,
            title="Управление ответами",
        )

        self.submission = Submission.objects.create(
            survey=self.survey,
            full_name="Участник",
            position="Технолог",
        )

    def test_submission_can_be_excluded(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse(
                "surveys:submission_exclude",
                args=[
                    self.survey.pk,
                    self.submission.pk,
                ],
            ),
            {
                "exclusion_reason": (
                    "Ответ отправлен ошибочно"
                ),
            },
        )

        self.submission.refresh_from_db()

        self.assertTrue(
            self.submission.is_excluded,
        )

        self.assertEqual(
            self.submission.exclusion_reason,
            "Ответ отправлен ошибочно",
        )

        self.assertEqual(
            response.status_code,
            302,
        )

    def test_submission_can_be_returned_to_results(self):
        self.submission.is_excluded = True
        self.submission.exclusion_reason = "Ошибка"
        self.submission.save()

        self.client.force_login(self.user)

        self.client.post(
            reverse(
                "surveys:submission_include",
                args=[
                    self.survey.pk,
                    self.submission.pk,
                ],
            ),
        )

        self.submission.refresh_from_db()

        self.assertFalse(
            self.submission.is_excluded,
        )

        self.assertEqual(
            self.submission.exclusion_reason,
            "",
        )

class SurveyProtocolTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="protocol_owner",
            password="StrongPassword123!",
        )

        self.survey = Survey.objects.create(
            owner=self.user,
            title="Протокольная дегустация",
        )

        self.sample = SurveySample.objects.create(
            survey=self.survey,
            name="Образец №1",
            description="Описание образца",
            order=10,
        )

        self.submission = Submission.objects.create(
            survey=self.survey,
            full_name="Иванов Иван Иванович",
            position="Инженер-технолог",
        )

        for question in self.survey.questions.all():
            if (
                question.question_type
                == SurveyQuestion.QuestionType.RATING
            ):
                Answer.objects.create(
                    submission=self.submission,
                    sample=self.sample,
                    question=question,
                    numeric_value=Decimal("4.5"),
                )
            else:
                Answer.objects.create(
                    submission=self.submission,
                    sample=self.sample,
                    question=question,
                    text_value="Хороший вкус",
                )

        self.survey.publish()
        self.survey.close()

        self.protocol = SurveyProtocol.objects.create(
            survey=self.survey,
            document_code="КК-Ф-020",
            protocol_number="304",
            protocol_date=date(2026, 5, 21),
            responsible_department="контроля качества",
            responsible_employee="Мальцева Е.С.",
            tasting_goal=("Органолептическая оценка продукции"),
            room_conditions=("Температура воздуха +18 – +25 °С."),
            product_conditions=("Температура продукции +55 ± 5 °С."),
            conclusion=("Образец получил высокие оценки."),
            signer_position="Нач. ОРП",
            signer_name="Холина О.В.",
        )

    def test_protocol_is_complete(self):
        self.assertTrue(
            self.protocol.is_complete,
        )

    def test_protocol_docx_is_created(self):
        document_data = build_protocol_docx(
            survey=self.survey,
            protocol=self.protocol,
        )

        self.assertTrue(
            document_data.startswith(b"PK"),
        )

        self.assertGreater(
            len(document_data),
            1000,
        )

    def test_protocol_page_available_after_close(self):
        self.client.force_login(self.user)

        response = self.client.get(
            reverse(
                "surveys:survey_protocol_edit",
                args=[
                    self.survey.pk,
                ],
            ),
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertContains(
            response,
            "Протокол дегустации",
        )

    def test_protocol_page_unavailable_for_published_survey(self):
        self.survey.status = Survey.Status.PUBLISHED
        self.survey.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        self.client.force_login(self.user)

        response = self.client.get(
            reverse(
                "surveys:survey_protocol_edit",
                args=[
                    self.survey.pk,
                ],
            ),
        )

        self.assertRedirects(
            response,
            reverse(
                "surveys:survey_results",
                args=[
                    self.survey.pk,
                ],
            ),
        )

    def test_protocol_date_is_written_to_docx(self):
        document_data = build_protocol_docx(
            survey=self.survey,
            protocol=self.protocol,
        )

        document = Document(BytesIO(document_data))

        document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        self.assertIn(
            "«21» мая 2026 г.",
            document_text,
        )

    def test_protocol_contains_conditions_section_titles(self):
        document_data = build_protocol_docx(
            survey=self.survey,
            protocol=self.protocol,
        )

        document = Document(BytesIO(document_data))

        document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        self.assertIn(
            "Условия в помещении:",
            document_text,
        )

        self.assertIn(
            "Температура и условия продуктов:",
            document_text,
        )

        self.assertIn(
            self.protocol.room_conditions,
            document_text,
        )

        self.assertIn(
            self.protocol.product_conditions,
            document_text,
        )

    def test_protocol_contains_responsible_department(self):
        document_data = build_protocol_docx(
            survey=self.survey,
            protocol=self.protocol,
        )

        document = Document(BytesIO(document_data))

        document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        self.assertIn(
            ("Сотрудник отдела контроля качества: Мальцева Е.С."),
            document_text,
        )


class GeneratedProtocolTests(TestCase):
    def setUp(self):
        self.temporary_media = tempfile.TemporaryDirectory()

        self.override = override_settings(
            MEDIA_ROOT=self.temporary_media.name,
        )

        self.override.enable()

        self.user = User.objects.create_user(
            username="generated_protocol_owner",
            password="StrongPassword123!",
        )

        self.survey = Survey.objects.create(
            owner=self.user,
            title="Форма для DOCX",
        )

        self.sample = SurveySample.objects.create(
            survey=self.survey,
            name="Образец №1",
            order=10,
        )

        submission = Submission.objects.create(
            survey=self.survey,
            full_name="Петров Пётр Петрович",
            position="Технолог",
        )

        for question in self.survey.questions.all():
            if (
                question.question_type
                == SurveyQuestion.QuestionType.RATING
            ):
                Answer.objects.create(
                    submission=submission,
                    sample=self.sample,
                    question=question,
                    numeric_value=Decimal("5.0"),
                )
            else:
                Answer.objects.create(
                    submission=submission,
                    sample=self.sample,
                    question=question,
                    text_value="Комментарий",
                )

        self.protocol = SurveyProtocol.objects.create(
            survey=self.survey,
            protocol_number="500",
            protocol_date=date(2026, 7, 28),
            responsible_employee="Сотрудник",
            tasting_goal="Оценка продукции",
            conclusion="Высокие оценки",
            signer_position="Начальник",
            signer_name="Иванов И.И.",
        )

    def tearDown(self):
        self.override.disable()
        self.temporary_media.cleanup()

    def test_new_versions_are_created(self):
        first = generate_and_save_protocol(
            survey=self.survey,
            protocol=self.protocol,
            user=self.user,
        )

        second = generate_and_save_protocol(
            survey=self.survey,
            protocol=self.protocol,
            user=self.user,
        )

        self.assertEqual(
            first.version,
            1,
        )

        self.assertEqual(
            second.version,
            2,
        )

        self.assertEqual(
            GeneratedProtocol.objects.count(),
            2,
        )

        self.assertTrue(
            first.file.name.endswith(
                ".docx"
            )
        )

        self.assertTrue(
            default_storage.exists(
                first.file.name,
            )
        )

class ProtocolDefaultsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="protocol_defaults_owner",
            password="StrongPassword123!",
        )

        self.previous_survey = Survey.objects.create(
            owner=self.user,
            title="Предыдущая форма",
        )

        self.previous_protocol = SurveyProtocol.objects.create(
            survey=self.previous_survey,
            document_code="КК-Ф-020",
            protocol_number="304",
            protocol_date=date(2026, 5, 21),
            responsible_department="контроля качества",
            responsible_employee="Мальцева Е.С.",
            tasting_goal="Органолептическая оценка",
            room_conditions="Температура +18 – +25 °С.",
            product_conditions="Продукт +55 ± 5 °С.",
            conclusion="Образцы получили высокие оценки.",
            signer_position="Нач. ОРП",
            signer_name="Холина О.В.",
        )

        self.new_survey = Survey.objects.create(
            owner=self.user,
            title="Новая форма",
        )

    def test_next_protocol_number(self):
        self.assertEqual(
            get_next_protocol_number(),
            "305",
        )

    def test_previous_fields_are_copied(self):
        initial_data = build_protocol_initial_data(
            survey=self.new_survey,
        )

        self.assertEqual(
            initial_data["protocol_number"],
            "305",
        )

        self.assertEqual(
            initial_data["responsible_department"],
            "контроля качества",
        )

        self.assertEqual(
            initial_data["responsible_employee"],
            "Мальцева Е.С.",
        )

        self.assertEqual(
            initial_data["tasting_goal"],
            "Органолептическая оценка",
        )

        self.assertEqual(
            initial_data["room_conditions"],
            "Температура +18 – +25 °С.",
        )

        self.assertEqual(
            initial_data["product_conditions"],
            "Продукт +55 ± 5 °С.",
        )

        self.assertEqual(
            initial_data["conclusion"],
            "Образцы получили высокие оценки.",
        )

        self.assertEqual(
            initial_data["signer_position"],
            "Нач. ОРП",
        )

        self.assertEqual(
            initial_data["signer_name"],
            "Холина О.В.",
        )

    def test_protocol_date_field_accepts_html_date(self):
        form = SurveyProtocolForm(
            data={
                "document_code": "КК-Ф-020",
                "protocol_number": "305",
                "protocol_date": "2026-07-28",
                "responsible_department": "разработки продуктов",
                "responsible_employee": "Сотрудник",
                "tasting_goal": "Цель",
                "room_conditions": "Условия",
                "product_conditions": "Температура",
                "conclusion": "Заключение",
                "signer_position": "Начальник",
                "signer_name": "Иванов И.И.",
            }
        )

        self.assertTrue(
            form.is_valid(),
            form.errors,
        )

        self.assertEqual(
            form.cleaned_data["protocol_date"],
            date(2026, 7, 28),
        )

    def test_protocol_date_is_rendered_for_html_input(self):
        protocol = SurveyProtocol.objects.create(
            survey=self.new_survey,
            protocol_number="305",
            protocol_date=date(2026, 7, 28),
        )

        form = SurveyProtocolForm(
            instance=protocol,
        )

        rendered_date = str(
            form["protocol_date"]
        )

        self.assertIn(
            'value="2026-07-28"',
            rendered_date,
        )


class SurveyCopyTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="copy_owner",
            password="StrongPassword123!",
        )

        self.source_survey = Survey.objects.create(
            owner=self.user,
            title="Исходная форма",
            description="Описание исходной формы",
            allow_multiple_submissions=True,
        )

        self.first_sample = SurveySample.objects.create(
            survey=self.source_survey,
            name="Образец №1",
            description="Первый образец",
            order=10,
        )

        self.second_sample = SurveySample.objects.create(
            survey=self.source_survey,
            name="Образец №2",
            description="Второй образец",
            order=20,
        )

    def test_survey_is_copied_as_draft(self):
        copied_survey = copy_survey(
            source_survey=self.source_survey,
            owner=self.user,
        )

        self.assertEqual(
            copied_survey.status,
            Survey.Status.DRAFT,
        )

        self.assertEqual(
            copied_survey.owner,
            self.user,
        )

        self.assertEqual(
            copied_survey.title,
            "Копия — Исходная форма",
        )

        self.assertNotEqual(
            copied_survey.public_id,
            self.source_survey.public_id,
        )

        self.assertIsNone(
            copied_survey.published_at,
        )

        self.assertIsNone(
            copied_survey.closed_at,
        )

    def test_samples_are_copied(self):
        copied_survey = copy_survey(
            source_survey=self.source_survey,
            owner=self.user,
        )

        copied_samples = list(
            copied_survey.samples.values_list(
                "name",
                "description",
                "order",
            )
        )

        self.assertEqual(
            copied_samples,
            [
                (
                    "Образец №1",
                    "Первый образец",
                    10,
                ),
                (
                    "Образец №2",
                    "Второй образец",
                    20,
                ),
            ],
        )

    def test_questions_are_copied_once(self):
        source_question_count = (
            self.source_survey.questions.count()
        )

        copied_survey = copy_survey(
            source_survey=self.source_survey,
            owner=self.user,
        )

        self.assertEqual(
            copied_survey.questions.count(),
            source_question_count,
        )

    def test_submissions_are_not_copied(self):
        Submission.objects.create(
            survey=self.source_survey,
            full_name="Участник",
            position="Технолог",
        )

        copied_survey = copy_survey(
            source_survey=self.source_survey,
            owner=self.user,
        )

        self.assertEqual(
            copied_survey.submissions.count(),
            0,
        )

    def test_protocol_is_not_copied(self):
        SurveyProtocol.objects.create(
            survey=self.source_survey,
            protocol_number="304",
            protocol_date=date(2026, 5, 21),
        )

        copied_survey = copy_survey(
            source_survey=self.source_survey,
            owner=self.user,
        )

        self.assertFalse(
            SurveyProtocol.objects.filter(
                survey=copied_survey,
            ).exists()
        )

    def test_copy_view_creates_form(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse(
                "surveys:survey_copy",
                args=[
                    self.source_survey.pk,
                ],
            ),
        )

        copied_survey = (
            Survey.objects
            .exclude(pk=self.source_survey.pk)
            .get()
        )

        self.assertRedirects(
            response,
            reverse(
                "surveys:survey_edit",
                args=[
                    copied_survey.pk,
                ],
            ),
        )

        self.assertEqual(
            copied_survey.status,
            Survey.Status.DRAFT,
        )


class SurveySectionListTests(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user(
            username="section_owner",
            password="StrongPassword123!",
        )

        self.other_user = User.objects.create_user(
            username="section_other",
            password="StrongPassword123!",
        )

        self.owner_survey = Survey.objects.create(
            owner=self.owner,
            title="Форма владельца",
        )

        self.other_survey = Survey.objects.create(
            owner=self.other_user,
            title="Чужая форма",
        )

        Submission.objects.create(
            survey=self.owner_survey,
            full_name="Участник владельца",
            position="Технолог",
        )

        Submission.objects.create(
            survey=self.other_survey,
            full_name="Чужой участник",
            position="Технолог",
        )

    def test_results_list_requires_login(self):
        response = self.client.get(
            reverse(
                "surveys:results_list",
            )
        )

        self.assertEqual(
            response.status_code,
            302,
        )

    def test_results_list_contains_own_survey(self):
        self.client.force_login(
            self.owner,
        )

        response = self.client.get(
            reverse(
                "surveys:results_list",
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertContains(
            response,
            "Форма владельца",
        )

    def test_results_list_hides_other_user_survey(self):
        self.client.force_login(
            self.owner,
        )

        response = self.client.get(
            reverse(
                "surveys:results_list",
            )
        )

        self.assertNotContains(
            response,
            "Чужая форма",
        )

    def test_results_list_hides_surveys_without_answers(self):
        empty_survey = Survey.objects.create(
            owner=self.owner,
            title="Форма без ответов",
        )

        self.client.force_login(
            self.owner,
        )

        response = self.client.get(
            reverse(
                "surveys:results_list",
            )
        )

        self.assertNotContains(
            response,
            empty_survey.title,
        )

    def test_protocol_list_requires_login(self):
        response = self.client.get(
            reverse(
                "surveys:protocol_list",
            )
        )

        self.assertEqual(
            response.status_code,
            302,
        )

    def test_protocol_list_contains_closed_own_survey(self):
        self.owner_survey.status = Survey.Status.CLOSED
        self.owner_survey.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        self.client.force_login(
            self.owner,
        )

        response = self.client.get(
            reverse(
                "surveys:protocol_list",
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertContains(
            response,
            "Форма владельца",
        )

    def test_protocol_list_hides_published_survey(self):
        self.owner_survey.status = Survey.Status.PUBLISHED
        self.owner_survey.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        self.client.force_login(
            self.owner,
        )

        response = self.client.get(
            reverse(
                "surveys:protocol_list",
            )
        )

        self.assertNotContains(
            response,
            "Форма владельца",
        )

    def test_protocol_list_hides_other_user_survey(self):
        self.owner_survey.status = Survey.Status.CLOSED
        self.owner_survey.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        self.other_survey.status = Survey.Status.CLOSED
        self.other_survey.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        self.client.force_login(
            self.owner,
        )

        response = self.client.get(
            reverse(
                "surveys:protocol_list",
            )
        )

        self.assertContains(
            response,
            "Форма владельца",
        )

        self.assertNotContains(
            response,
            "Чужая форма",
        )

    def test_protocol_list_shows_protocol_details(self):
        self.owner_survey.status = Survey.Status.CLOSED
        self.owner_survey.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        SurveyProtocol.objects.create(
            survey=self.owner_survey,
            protocol_number="501",
            protocol_date=date(
                2026,
                7,
                30,
            ),
        )

        self.client.force_login(
            self.owner,
        )

        response = self.client.get(
            reverse(
                "surveys:protocol_list",
            )
        )

        self.assertContains(
            response,
            "№ 501",
        )

        self.assertContains(
            response,
            "30.07.2026",
        )


class SurveyPaginationTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="pagination_user",
            password="StrongPassword123!",
        )

        self.client.force_login(
            self.user,
        )

    def create_surveys(
        self,
        count,
        *,
        status=Survey.Status.DRAFT,
        with_submissions=False,
    ):
        surveys = []

        for index in range(count):
            survey = Survey.objects.create(
                owner=self.user,
                title=f"Форма {index + 1:02d}",
                status=status,
            )

            if with_submissions:
                Submission.objects.create(
                    survey=survey,
                    full_name=f"Участник {index + 1}",
                    position="Технолог",
                )

            surveys.append(survey)

        return surveys

    def test_survey_list_is_paginated(self):
        self.create_surveys(21)

        response = self.client.get(
            reverse(
                "surveys:survey_list",
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertEqual(
            len(response.context["surveys"]),
            20,
        )

        self.assertEqual(
            response.context["page_obj"].paginator.count,
            21,
        )

        self.assertEqual(
            response.context["page_obj"].paginator.num_pages,
            2,
        )

    def test_survey_list_second_page(self):
        self.create_surveys(21)

        response = self.client.get(
            reverse(
                "surveys:survey_list",
            ),
            {
                "page": 2,
            },
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertEqual(
            len(response.context["surveys"]),
            1,
        )

        self.assertEqual(
            response.context["page_obj"].number,
            2,
        )

    def test_results_list_is_paginated(self):
        self.create_surveys(
            21,
            with_submissions=True,
        )

        response = self.client.get(
            reverse(
                "surveys:results_list",
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertEqual(
            len(response.context["surveys"]),
            20,
        )

        self.assertEqual(
            response.context["page_obj"].paginator.count,
            21,
        )

    def test_protocol_list_is_paginated(self):
        self.create_surveys(
            21,
            status=Survey.Status.CLOSED,
        )

        response = self.client.get(
            reverse(
                "surveys:protocol_list",
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertEqual(
            len(response.context["survey_rows"]),
            20,
        )

        self.assertEqual(
            response.context["page_obj"].paginator.count,
            21,
        )

    def test_invalid_page_falls_back_safely(self):
        self.create_surveys(21)

        response = self.client.get(
            reverse(
                "surveys:survey_list",
            ),
            {
                "page": "invalid",
            },
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertEqual(
            response.context["page_obj"].number,
            1,
        )


class SurveyLiveSearchTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="search_user",
            password="StrongPassword123!",
            first_name="Иван",
            last_name="Иванов",
        )

        self.other_user = User.objects.create_user(
            username="other_search_user",
            password="StrongPassword123!",
            first_name="Пётр",
            last_name="Петров",
        )

        self.client.force_login(
            self.user,
        )

        self.milk_survey = Survey.objects.create(
            owner=self.user,
            title="Дегустация молока",
            description="Сравнение образцов молочной продукции",
        )

        self.juice_survey = Survey.objects.create(
            owner=self.user,
            title="Дегустация сока",
            description="Апельсиновый сок",
        )

        self.other_survey = Survey.objects.create(
            owner=self.other_user,
            title="Чужая форма молока",
            description="Чужая форма",
        )

    def test_survey_list_searches_by_title(self):
        response = self.client.get(
            reverse(
                "surveys:survey_list",
            ),
            {
                "q": "молока",
            },
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertContains(
            response,
            "Дегустация молока",
        )

        self.assertNotContains(
            response,
            "Дегустация сока",
        )

    def test_survey_list_search_is_case_insensitive(self):
        response = self.client.get(
            reverse(
                "surveys:survey_list",
            ),
            {
                "q": "МОЛОКА",
            },
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertContains(
            response,
            "Дегустация молока",
        )

    def test_survey_list_searches_by_description(self):
        response = self.client.get(
            reverse(
                "surveys:survey_list",
            ),
            {
                "q": "апельсиновый",
            },
        )

        self.assertContains(
            response,
            "Дегустация сока",
        )

        self.assertNotContains(
            response,
            "Дегустация молока",
        )

    def test_survey_list_does_not_show_other_user_surveys(self):
        response = self.client.get(
            reverse(
                "surveys:survey_list",
            ),
            {
                "q": "молока",
            },
        )

        self.assertContains(
            response,
            "Дегустация молока",
        )

        self.assertNotContains(
            response,
            "Чужая форма молока",
        )

    def test_results_search(self):
        Submission.objects.create(
            survey=self.milk_survey,
            full_name="Участник молока",
            position="Технолог",
        )

        Submission.objects.create(
            survey=self.juice_survey,
            full_name="Участник сока",
            position="Технолог",
        )

        response = self.client.get(
            reverse(
                "surveys:results_list",
            ),
            {
                "q": "молока",
            },
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertContains(
            response,
            "Дегустация молока",
        )

        self.assertNotContains(
            response,
            "Дегустация сока",
        )

    def test_protocol_searches_by_title(self):
        self.milk_survey.status = Survey.Status.CLOSED
        self.milk_survey.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        self.juice_survey.status = Survey.Status.CLOSED
        self.juice_survey.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        response = self.client.get(
            reverse(
                "surveys:protocol_list",
            ),
            {
                "q": "молока",
            },
        )

        self.assertContains(
            response,
            "Дегустация молока",
        )

        self.assertNotContains(
            response,
            "Дегустация сока",
        )

    def test_protocol_searches_by_protocol_number(self):
        self.milk_survey.status = Survey.Status.CLOSED
        self.milk_survey.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        SurveyProtocol.objects.create(
            survey=self.milk_survey,
            protocol_number="501",
        )

        response = self.client.get(
            reverse(
                "surveys:protocol_list",
            ),
            {
                "q": "501",
            },
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertContains(
            response,
            "Дегустация молока",
        )

        self.assertContains(
            response,
            "501",
        )

    def test_search_query_is_preserved_in_context(self):
        response = self.client.get(
            reverse(
                "surveys:survey_list",
            ),
            {
                "q": "молока",
            },
        )

        self.assertEqual(
            response.context["search_query"],
            "молока",
        )

    def test_empty_query_returns_normal_list(self):
        response = self.client.get(
            reverse(
                "surveys:survey_list",
            ),
            {
                "q": "",
            },
        )

        self.assertContains(
            response,
            "Дегустация молока",
        )

        self.assertContains(
            response,
            "Дегустация сока",
        )

    def test_search_is_paginated(self):
        for index in range(25):
            Survey.objects.create(
                owner=self.user,
                title=f"Поисковая форма {index + 1}",
                description="Специальный поиск",
            )

        response = self.client.get(
            reverse(
                "surveys:survey_list",
            ),
            {
                "q": "Поисковая форма",
            },
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertEqual(
            response.context["page_obj"].paginator.count,
            25,
        )

        self.assertEqual(
            response.context["page_obj"].paginator.num_pages,
            2,
        )

        self.assertEqual(
            len(response.context["surveys"]),
            20,
        )

    def test_search_second_page(self):
        for index in range(25):
            Survey.objects.create(
                owner=self.user,
                title=f"Поисковая форма {index + 1}",
                description="Специальный поиск",
            )

        response = self.client.get(
            reverse(
                "surveys:survey_list",
            ),
            {
                "q": "Поисковая форма",
                "page": 2,
            },
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertEqual(
            response.context["page_obj"].number,
            2,
        )

        self.assertEqual(
            len(response.context["surveys"]),
            5,
        )

    def test_ajax_search_returns_success(self):
        response = self.client.get(
            reverse(
                "surveys:survey_list",
            ),
            {
                "q": "молока",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertContains(
            response,
            "Дегустация молока",
        )

