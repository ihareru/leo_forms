from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Side,
    Border,
    )
from openpyxl.utils import get_column_letter

from apps.surveys.models import SurveyQuestion
from apps.surveys.services.results import get_survey_results
from django.utils.timezone import localtime


HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7",
)

THIN_SIDE = Side(
    style="thin",
    color="B7B7B7",
)

CELL_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)


def configure_cell(
    cell,
    *,
    bold=False,
    wrap_text=True,
    horizontal="left",
    vertical="top",
):
    cell.font = Font(
        bold=bold,
    )

    cell.alignment = Alignment(
        horizontal=horizontal,
        vertical=vertical,
        wrap_text=wrap_text,
    )

    cell.border = CELL_BORDER


def set_column_widths(
    worksheet,
    widths,
):
    for column_number, width in widths.items():
        column_letter = get_column_letter(column_number)
        worksheet.column_dimensions[column_letter].width = width


def create_summary_sheet(
    workbook,
    survey,
):
    results = get_survey_results(survey)

    worksheet = workbook.active
    worksheet.title = "Сводные результаты"

    worksheet.append(
        [
            "Форма",
            survey.title,
        ]
    )

    worksheet.append(
        [
            "Всего ответов",
            results["submissions_total"],
        ]
    )

    worksheet.append(
        [
            "Участвуют в расчётах",
            results["submissions_included"],
        ]
    )

    worksheet.append(
        [
            "Исключено",
            results["submissions_excluded"],
        ]
    )

    worksheet.append([])

    headers = [
        "№",
        "Образец",
        "Характеристика",
    ]

    headers.extend(
        question.title
        for question in results["rating_questions"]
    )

    headers.extend(
        [
            "Средний балл",
            "Комментарии",
        ]
    )

    worksheet.append(headers)

    header_row_number = worksheet.max_row

    for cell in worksheet[header_row_number]:
        configure_cell(
            cell,
            bold=True,
            horizontal="center",
            vertical="center",
        )

        cell.fill = HEADER_FILL

    for row_number, row_data in enumerate(
        results["rows"],
        start=1,
    ):
        comments_text = "\n".join(
            row_data["comments"]
        )

        row = [
            row_number,
            row_data["sample"].name,
            row_data["sample"].description,
        ]

        row.extend(
            score["average"]
            for score in row_data["scores"]
        )

        row.extend(
            [
                row_data["overall_average"],
                comments_text,
            ]
        )

        worksheet.append(row)

    for row in worksheet.iter_rows(
        min_row=header_row_number + 1,
    ):
        for cell in row:
            configure_cell(cell)

    worksheet.freeze_panes = (
        f"A{header_row_number + 1}"
    )

    worksheet.auto_filter.ref = (
        f"A{header_row_number}:"
        f"{get_column_letter(worksheet.max_column)}"
        f"{worksheet.max_row}"
    )

    set_column_widths(
        worksheet,
        {
            1: 6,
            2: 25,
            3: 55,
        },
    )

    for column_number in range(
        4,
        4 + len(results["rating_questions"]),
    ):
        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = 15

    average_column = (
        4 + len(results["rating_questions"])
    )

    comments_column = average_column + 1

    worksheet.column_dimensions[
        get_column_letter(average_column)
    ].width = 15

    worksheet.column_dimensions[
        get_column_letter(comments_column)
    ].width = 60


def create_detailed_answers_sheet(
    workbook,
    survey,
):
    worksheet = workbook.create_sheet(
        title="Ответы участников",
    )

    samples = list(
        survey.samples.order_by(
            "order",
            "id",
        )
    )

    questions = list(
        survey.questions.order_by(
            "order",
            "id",
        )
    )

    headers = [
        "№",
        "ФИО",
        "Должность",
        "Дата и время",
        "Исключён из расчётов",
        "Причина исключения",
    ]

    answer_columns = []

    for sample in samples:
        for question in questions:
            headers.append(
                f"{sample.name}: {question.title}"
            )

            answer_columns.append(
                (
                    sample,
                    question,
                )
            )

    worksheet.append(headers)

    for cell in worksheet[1]:
        configure_cell(
            cell,
            bold=True,
            horizontal="center",
            vertical="center",
        )

        cell.fill = HEADER_FILL

    submissions = (
        survey.submissions
        .prefetch_related(
            "answers",
            "answers__sample",
            "answers__question",
        )
        .order_by(
            "submitted_at",
            "id",
        )
    )

    for row_number, submission in enumerate(
        submissions,
        start=1,
    ):
        answer_map = {
            (
                answer.sample_id,
                answer.question_id,
            ): answer
            for answer in submission.answers.all()
        }

        row = [
            row_number,
            submission.full_name,
            submission.position,
            localtime(
                submission.submitted_at,
            ).strftime(
                "%d.%m.%Y %H:%M:%S"
            ),
            "Да" if submission.is_excluded else "Нет",
            submission.exclusion_reason,
        ]

        for sample, question in answer_columns:
            answer = answer_map.get(
                (
                    sample.pk,
                    question.pk,
                )
            )

            if answer is None:
                row.append("")
                continue

            if (
                question.question_type
                == SurveyQuestion.QuestionType.RATING
            ):
                row.append(answer.numeric_value)
            else:
                row.append(answer.text_value)

        worksheet.append(row)

    for row in worksheet.iter_rows(
        min_row=2,
    ):
        for cell in row:
            configure_cell(cell)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:"
        f"{get_column_letter(worksheet.max_column)}"
        f"{worksheet.max_row}"
    )

    set_column_widths(
        worksheet,
        {
            1: 6,
            2: 32,
            3: 28,
            4: 22,
            5: 20,
            6: 40,
        },
    )

    for column_number in range(
        7,
        worksheet.max_column + 1,
    ):
        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = 28


def build_survey_excel_report(
    survey,
) -> bytes:
    """
    Формирует Excel-файл со сводными и именными данными.
    """

    workbook = Workbook()

    create_summary_sheet(
        workbook,
        survey,
    )

    create_detailed_answers_sheet(
        workbook,
        survey,
    )

    buffer = BytesIO()

    workbook.save(buffer)

    return buffer.getvalue()